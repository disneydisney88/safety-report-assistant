from __future__ import annotations

from io import BytesIO
import re
from typing import BinaryIO

from docx import Document
from pypdf import PdfReader


MAX_EXTRACT_CHARS = 18000


def _read_bytes(file_obj: BinaryIO | BytesIO | bytes) -> bytes:
    if isinstance(file_obj, bytes):
        return file_obj
    if hasattr(file_obj, "getvalue"):
        return file_obj.getvalue()
    return file_obj.read()


def extract_text_from_upload(file_obj: BinaryIO | BytesIO | bytes, file_name: str) -> str:
    suffix = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    raw = _read_bytes(file_obj)
    if suffix == "docx":
        return extract_docx_text(raw)
    if suffix == "pdf":
        return extract_pdf_text(raw)
    if suffix == "txt":
        return raw.decode("utf-8", errors="ignore")[:MAX_EXTRACT_CHARS]
    raise ValueError("Only .docx, .pdf and .txt Method Statement files are supported.")


def extract_docx_text(raw: bytes) -> str:
    doc = Document(BytesIO(raw))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " | ") for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)[:MAX_EXTRACT_CHARS]


def _ascii_space_health(text: str) -> float:
    """Spaces per ASCII letter — near-zero means the extractor dropped spaces.

    Normal English runs at roughly one space per 5-6 letters (~0.17); broken
    extractions like "Bypassbutton(bypassslackropelimit" sit near zero.
    """
    letters = sum(1 for ch in text if ch.isascii() and ch.isalpha())
    if letters < 200:
        return 1.0
    return text.count(" ") / letters


def _pypdf_pages(raw: bytes, layout: bool) -> str:
    reader = PdfReader(BytesIO(raw))
    parts: list[str] = []
    for page in reader.pages[:40]:
        if layout:
            try:
                text = page.extract_text(extraction_mode="layout") or ""
            except Exception:
                text = ""
        else:
            text = page.extract_text() or ""
        text = text.strip()
        if text:
            parts.append(text)
    return "\n".join(parts)


def _pdfplumber_pages(raw: bytes) -> str:
    try:
        import pdfplumber
    except Exception:
        return ""
    parts: list[str] = []
    try:
        with pdfplumber.open(BytesIO(raw)) as pdf:
            for page in pdf.pages[:40]:
                text = (page.extract_text() or "").strip()
                if text:
                    parts.append(text)
    except Exception:
        return ""
    return "\n".join(parts)


def extract_pdf_text(raw: bytes) -> str:
    # Some PDFs (certain CAD/Word exports) lose inter-word spaces under pypdf's
    # default extractor. Escalate: pypdf -> pypdf layout mode -> pdfplumber ->
    # heuristic word-splitting repair, keeping the healthiest result.
    extracted = _pypdf_pages(raw, layout=False)
    if _ascii_space_health(extracted) < 0.05:
        candidates = [extracted, _pypdf_pages(raw, layout=True), _pdfplumber_pages(raw)]
        extracted = max(candidates, key=lambda t: (_ascii_space_health(t), len(t)))
        if _ascii_space_health(extracted) < 0.05:
            extracted = repair_concatenated_english(extracted)
    extracted = extracted[:MAX_EXTRACT_CHARS]
    if not extracted.strip():
        raise ValueError("No selectable text found in this PDF. It may be a scanned PDF; OCR support is not enabled yet.")
    return extracted


def _compact_text(value: str) -> str:
    return "".join(str(value or "").split()).lower()


# Vocabulary for repairing space-less PDF extractions. Ordered lookup is by
# length (longest match first) so "structural" wins over "structure"+junk.
# Construction / BMU / scaffolding / E&M terms plus the common glue words that
# appear in method statements.
_REPAIR_VOCAB = {
    # BMU / gondola / T&C
    "bypass", "slack", "rope", "ropes", "limit", "limits", "cage", "cradle",
    "gondola", "hoist", "hoists", "trolley", "jib", "jibhead", "telescopic",
    "slewing", "traversing", "traverser", "winch", "davit", "counterweight",
    "suspended", "platform", "anchorage", "wire", "swl", "bmu",
    # generic construction / inspection
    "visual", "check", "checks", "checked", "checking", "inspect", "inspection",
    "test", "tests", "testing", "load", "loads", "loading", "unload", "structural",
    "structure", "structures", "members", "member", "connection", "connections",
    "bolts", "bolt", "nuts", "nut", "nets", "net", "screws", "washer", "bracket",
    "brackets", "rail", "rails", "track", "tracks", "wheel", "wheels", "wheelsets",
    "scaffold", "scaffolding", "bamboo", "formwork", "rebar", "concrete",
    # actions
    "install", "installation", "erect", "erection", "remove", "removal",
    "dismantle", "dismantling", "lift", "lifting", "lower", "lowering", "raise",
    "operate", "operation", "operations", "connect", "connection", "energize",
    "commission", "commissioning", "witness", "verify", "confirm", "ensure",
    "carry", "conduct", "perform", "record", "sign", "issue", "apply", "land",
    "landing", "stop", "start", "press", "push", "pull", "turn", "switch",
    "activate", "deactivate", "restore", "reset",
    # equipment / electrical
    "power", "supply", "cable", "cables", "voltage", "phase", "phases",
    "isolator", "breaker", "elcb", "rcd", "socket", "plug", "motor", "button",
    "buttons", "emergency", "alarm", "brake", "brakes", "device", "devices",
    "controller", "control", "controls", "panel", "key", "selector",
    # safety / documents
    "safety", "harness", "lanyard", "helmet", "permit", "permits", "form",
    "certificate", "certificates", "competent", "person", "examiner", "operator",
    "engineer", "supervisor", "worker", "workers", "public", "barrier",
    "barriers", "warning", "signage", "exclusion", "zone", "rescue", "weather",
    # glue words
    "the", "a", "an", "and", "or", "of", "to", "in", "on", "at", "by", "for",
    "with", "from", "all", "any", "each", "before", "after", "during", "under",
    "over", "up", "down", "left", "right", "can", "shall", "must", "will", "be",
    "is", "are", "not", "no", "so", "as", "if", "then", "when", "upper", "lower",
    "ground", "roof", "site", "area", "work", "works", "working", "step", "steps",
    "procedure", "procedures", "method", "statement", "material", "materials",
    "obstruction", "bar", "buffer", "end", "travel", "speed", "direction",
    "function", "functions", "static", "dynamic", "weight", "weights", "kg",
    "mm", "m", "min", "max", "high", "level", "levels",
}
_MAX_VOCAB_LEN = max(len(word) for word in _REPAIR_VOCAB)
# Two-letter words allowed in dictionary segmentation. Longer minimum match
# elsewhere avoids spurious splits, but these appear constantly in MS text
# ("so cage can land ON ground", "UP/DOWN").
_SHORT_GLUE_WORDS = {"of", "to", "in", "on", "at", "by", "up", "so", "no", "if", "as", "be", "is", "or", "an"}


def _split_lowercase_run(run: str) -> str:
    """Best-effort dictionary segmentation of a space-less lowercase run.

    Dynamic programme minimising unmatched characters, preferring longer
    dictionary words. Unmatched remainders are kept verbatim so technical
    tokens (e.g. product codes) survive unchanged.
    """
    n = len(run)
    lowered = run.lower()
    # cost[i] = (unmatched_chars, word_count) for best split of run[:i]
    cost: list[tuple[int, int]] = [(0, 0)] + [(n + 1, 0)] * n
    choice: list[tuple[int, bool] | None] = [None] * (n + 1)
    for i in range(1, n + 1):
        # Treat run[j:i] as one unmatched char appended to run[:i-1].
        best = (cost[i - 1][0] + 1, cost[i - 1][1] + 1)
        best_choice: tuple[int, bool] = (i - 1, False)
        for length in range(min(_MAX_VOCAB_LEN, i), 1, -1):
            j = i - length
            word = lowered[j:i]
            if word in _REPAIR_VOCAB and (length >= 3 or word in _SHORT_GLUE_WORDS):
                cand = (cost[j][0], cost[j][1] + 1)
                if cand < best:
                    best = cand
                    best_choice = (j, True)
        cost[i] = best
        choice[i] = best_choice
    # Short runs may be single legitimate words ("witnessed") — only split
    # them on a perfect segmentation into 2+ dictionary words. Long runs are
    # clearly concatenated, so a partial split beats leaving them glued.
    unmatched, words = cost[n]
    if n < 16:
        if unmatched > 0 or words < 2:
            return run
    elif unmatched > n // 3:
        return run
    parts: list[str] = []
    i = n
    pending = ""
    while i > 0:
        j, matched = choice[i]  # type: ignore[misc]
        if matched:
            if pending:
                parts.append(pending)
                pending = ""
            parts.append(run[j:i])
        else:
            pending = run[j:i] + pending
        i = j
    if pending:
        parts.append(pending)
    parts.reverse()
    return " ".join(parts)


def repair_concatenated_english(text: str) -> str:
    """Repair PDF text whose inter-word spaces were lost during extraction."""
    # Space after sentence punctuation stuck to the next word.
    text = re.sub(r"([,.;:!?])([A-Za-z])", r"\1 \2", text)
    # Space around brackets glued to words.
    text = re.sub(r"([A-Za-z0-9])\(", r"\1 (", text)
    text = re.sub(r"\)([A-Za-z])", r") \1", text)
    # camelCase boundary from merged words, e.g. "Visualcheck" stays but
    # "checkAll" splits; also digit-to-letter like "150%SWL" -> "150% SWL".
    text = re.sub(r"([a-z])([A-Z][a-z])", r"\1 \2", text)
    # Lowercase glued to an acronym, e.g. "ofBMU" -> "of BMU".
    text = re.sub(r"([a-z])([A-Z]{2,})", r"\1 \2", text)
    text = re.sub(r"(%|\d)([A-Za-z]{3,})", r"\1 \2", text)
    # Word glued to a number, e.g. "Connect380V" (lowercase-before-digit only,
    # so codes like "M12" survive). Also ",3phases" -> ", 3phases".
    text = re.sub(r"([a-z])(\d)", r"\1 \2", text)
    text = re.sub(r"([,;])(\d)", r"\1 \2", text)
    # Dictionary-split long letter runs such as "bypassslackropelimit" or
    # "Bypassbutton" (case-insensitive segmentation, original case preserved).
    def _fix(match: re.Match[str]) -> str:
        return _split_lowercase_run(match.group(0))

    return re.sub(r"[A-Za-z][a-z]{6,}", _fix, text)


# Chinese and English action verbs that mark a real physical work step.
# Includes testing & commissioning verbs (測試/檢查/確認/接駁) so T&C method
# statements such as BMU / gondola testing keep their real steps.
STEP_ACTION_TERMS = [
    "拆", "安裝", "搭", "吊", "搬", "運", "傳", "清", "封閉", "移除", "鋪", "綁", "紮",
    "澆", "挖", "焊", "切割", "鑽", "測試", "檢查", "確認", "接駁", "調試", "啟動",
    "set up", "install", "erect", "remove", "dismantle",
    "transport", "carry out", "lift", "hoist", "lower", "pour", "excavate", "weld",
    "cut", "drill", "deliver", "position", "inspect", "test", "check", "connect",
    "energize", "commission", "verify", "witness",
]

# Suffixes that mark a document title or section heading when the line is short
# and has no action verb, e.g. 拆棚施工方案 / 安全程序及措施 / 工地要求.
_HEADING_SUFFIXES = ("方案", "方法書", "程序", "措施", "安排", "要求", "守則", "規定", "清單", "注意事項", "工序", "圖解", "示意圖")

# Lines starting with these are safety rules / control measures, not steps.
_CONTROL_STARTERS = ("必須", "嚴禁", "不得", "切勿", "確保", "所有工人", "工人須", "工人必須", "如遇", "如有", "如因", "為免", "為防")


def clean_extracted_steps(raw_steps: list[str], titles: list[str] | None = None, max_steps: int = 14) -> list[str]:
    """Keep only true physical work steps; drop titles, headings and controls."""
    title_keys = {_compact_text(title) for title in (titles or []) if title}
    always_blocked_terms = [
        "施工方案",
        "methodstatement",
        "安全程序及措施",
        "拆棚之程序",
        "riskassessment",
        "tableofcontents",
        # Table headers / step-list captions, e.g. "步驟 中文工序".
        "步驟中文工序",
        "中文工序",
        "工序標題",
        "簡化版工序",
        # Illustrated step-guide titles, e.g. "簡易中文施工步驟圖解".
        "施工步驟圖解",
        "步驟圖解",
    ]
    blocked_terms = [
        "安全準備",
        "準備工作",
        "適用法例",
        "工地要求",
        "拆棚後安排",
        "個人防護",
        "ppe",
        "訓練",
        "permit",
        "許可證",
        "必須",
        "嚴禁",
        "如遇天氣",
        "惡劣天氣",
        "toolbox",
        "工具箱",
    ]
    steps: list[str] = []
    for raw in raw_steps:
        clean = str(raw or "").strip(" \t-*0123456789.)、")
        # Repair steps carried over from a space-less PDF extraction, e.g.
        # "Bypassbutton(bypassslackropelimit,..."; drop them if unrepairable —
        # broken strings must never reach a formal RA.
        if re.search(r"[A-Za-z]{15,}", clean):
            clean = repair_concatenated_english(clean)
            if re.search(r"[A-Za-z]{18,}", clean):
                continue
            # A repaired heading ("Method of Site Testing of BMU") reads like
            # a step but is still a document/section title, not a work step.
            if _compact_text(clean).startswith(("methodof", "procedureof", "t&cprocedure", "tableof", "appendix")):
                continue
        compact = _compact_text(clean)
        if not clean or len(compact) < 6:
            continue
        if compact in title_keys or any(compact == title or compact in title for title in title_keys):
            continue
        if any(term in compact for term in always_blocked_terms):
            continue
        # Concatenated English fragments from drawings / table cells such as
        # "Cageupbutton" or "T&CProcedureof" are labels, not work steps.
        if re.fullmatch(r"[A-Za-z&()/\-.0-9]+", clean) and len(clean) >= 8:
            continue
        # Component labels (buttons, switches, selectors) are not steps.
        if clean.lower().rstrip(" )").endswith(("button", "switch", "selector", "indicator", "checklist")):
            continue
        has_action = any(term in clean.lower() for term in STEP_ACTION_TERMS)
        # Short heading-like lines without an action verb are titles/headings.
        if not has_action and len(compact) <= 32 and clean.endswith(_HEADING_SUFFIXES):
            continue
        # Safety rules / control measures are not steps.
        if clean.startswith(_CONTROL_STARTERS):
            continue
        if any(term in compact for term in blocked_terms) and not has_action:
            continue
        if clean not in steps:
            steps.append(clean[:260])
        if len(steps) >= max_steps:
            break
    return steps


def infer_steps_from_ms_text(text: str, max_steps: int = 14) -> list[str]:
    def normalize_step(value: str) -> str:
        clean = re.sub(r"\s+", "", value.strip())
        clean = re.sub(r"^[一二三四五六七八九十\d]+[\.、:：)\)]", "", clean)
        clean = re.sub(r"(.{2,12})\1{2,}", r"\1", clean)
        return clean.strip(" ：:.-")

    def normalize_line(value: str) -> str:
        clean = re.sub(r"\s+", " ", value.strip())
        clean = re.sub(r"^\s*\d+\s*[\.)、]?\s*", "", clean)
        clean = clean.replace(" /", "/").replace("/ ", "/")
        clean = clean.replace(" 。", "。").replace(" ，", "，")
        clean = re.sub(r"(.{6,20})\1{1,}", r"\1", clean)
        return clean.strip(" ：:.-")

    def is_heading_or_control(value: str) -> bool:
        clean = normalize_step(value)
        if len(clean) < 6 or len(clean) > 240:
            return True
        if re.search(r"(.{3,12})\1{1,}", clean):
            return True
        heading_terms = ["施工方案", "安全程序及措施", "安全準備", "準備工作", "拆棚後安排", "適用法例", "工地要求", "圖一", "圖二"]
        if any(term in clean for term in heading_terms) and len(clean) <= 32:
            return True
        control_starts = ["必須", "所有", "每日", "如遇", "如因", "如竹棚", "為免", "嚴禁", "參加", "合資格", "派遣", "安裝足夠", "有關實際"]
        if any(clean.startswith(term) for term in control_starts):
            return True
        control_terms = ["大原則", "訓練", "佩帶", "安全帽", "安全帶", "嚴禁", "檢查", "記錄", "提醒工人", "非作業人員"]
        # Broad list: a line naming real construction activities is a step even
        # when it also mentions 檢查/記錄 (e.g. 步驟10 新造樓板及樑 ... 經檢查並
        # 簽發 T4 後才可澆築混凝土 was wrongly dropped as a control).
        action_terms = [
            "拆除", "拆卸", "傳遞", "運走", "清走", "安裝", "封閉",
            "植筋", "扎鐵", "紮鐵", "澆築", "澆注", "落石屎", "混凝土工程",
            "搭設", "搭建", "鑽孔", "鑽切", "切割", "吊起", "吊運", "回頂",
            "圍封", "模板", "支架工程", "接駁", "測試", "調試",
        ]
        if any(term in clean for term in control_terms) and not any(term in clean for term in action_terms):
            return True
        return False

    def unique_append(target: list[str], value: str) -> None:
        clean = normalize_line(value)
        compact = normalize_step(clean)
        if clean and clean not in target and not is_heading_or_control(compact):
            target.append(clean[:260])

    lines = (text or "").splitlines()

    # Prefer the real construction sequence section, e.g. "二. 拆棚工序".
    in_sequence = False
    section_steps: list[str] = []
    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        compact = normalize_step(line)
        idx += 1
        if not compact:
            continue
        if any(marker in compact for marker in ["拆棚工序", "施工工序", "施工程序", "工作程序", "工作步驟"]):
            in_sequence = True
            continue
        if in_sequence and re.match(r"^[ABC]\)?|^[一二三四五六七八九十][\.:：]", line) and section_steps:
            break
        if in_sequence and re.match(r"^\s*\d+[\.)、]\s*", line):
            combined = line
            while idx < len(lines):
                next_line = lines[idx].strip()
                if not next_line:
                    idx += 1
                    continue
                if re.match(r"^\s*\d+[\.)、]\s*", next_line) or re.match(r"^[ABC]\)?|^[一二三四五六七八九十][\.:：]", next_line):
                    break
                compact_next = normalize_step(next_line)
                if (
                    any(stop in compact_next for stop in ["有限公司", "Tel", "Fax", "____", "WINGLEE", "SCAFFOLDING", "HennessyRoad"])
                    or re.fullmatch(r"\d+/\d+", compact_next)
                ):
                    idx += 1
                    continue
                next_line = re.sub(r"[（(]\s*圖[一二三四五六七八九十\d]+\s*[）)]", "", next_line).strip()
                if not next_line:
                    idx += 1
                    continue
                combined += next_line
                idx += 1
            unique_append(section_steps, combined)
            if len(section_steps) >= max_steps:
                return section_steps
    if section_steps:
        return section_steps[:max_steps]

    steps: list[str] = []
    skip_first_cells = {
        "work activity",
        "project / area",
        "severity (s)",
        "likelihood (l)",
        "risk score",
        "prepared by",
    }
    for line in lines:
        cells = [cell.strip() for cell in line.split(" | ")]
        if len(cells) < 5:
            continue
        first_cell = normalize_step(cells[0])
        if not first_cell or first_cell.lower() in skip_first_cells:
            continue
        unique_append(steps, first_cell)
        if len(steps) >= max_steps:
            return steps

    markers = ("step", "sequence", "procedure", "work sequence", "工序", "程序", "步驟")
    for line in lines:
        clean = normalize_step(line)
        if not clean or len(clean) < 8:
            continue
        lower = clean.lower()
        if lower.startswith(("method statement", "risk assessment", "table of contents")):
            continue
        numbered_line = bool(re.match(r"^\s*(\d+[\.)、]|step\s+\d+)", line, flags=re.IGNORECASE))
        looks_like_step = numbered_line or any(marker in lower for marker in markers)
        if looks_like_step:
            unique_append(steps, clean)
        if len(steps) >= max_steps:
            break
    return steps
